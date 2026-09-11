"""Preserve all attached source material for the offline atlas (no network)."""
from pathlib import Path
import json, base64, zipfile, re, xml.etree.ElementTree as ET, hashlib
from docx import Document
from docx.text.paragraph import Paragraph
from docx.table import Table
from docx.oxml.ns import qn
ROOT=Path(__file__).resolve().parents[1]
P=ROOT/'originals'
if not P.exists(): P=Path('/mnt/data')
OUT=ROOT/'data'
OUT.mkdir(exist_ok=True)

def docx_blocks(path):
    doc=Document(path); blocks=[]; sections=[]; sec={'id':'introduction','title':'Cover and executive summary','blocks':[]}
    sections.append(sec); n=0
    for child in doc.element.body:
        if child.tag==qn('w:p'):
            p=Paragraph(child,doc); text=p.text
            if text.strip():
                if p.style and p.style.name.startswith('Heading'):
                    n+=1; sec={'id':f'section-{n}','title':text,'blocks':[]}; sections.append(sec)
                sec['blocks'].append({'type':'paragraph','style':p.style.name if p.style else 'Normal','text':text})
            for img in child.findall('.//'+qn('a:blip')):
                rid=img.get(qn('r:embed'))
                if rid in doc.part.related_parts:
                    part=doc.part.related_parts[rid]; mime=part.content_type
                    sec['blocks'].append({'type':'image','data':f'data:{mime};base64,'+base64.b64encode(part.blob).decode(),'caption':'Original report figure — archived methodology; see corrections.'})
        elif child.tag==qn('w:tbl'):
            t=Table(child,doc)
            rows=[[c.text for c in row.cells] for row in t.rows]
            sec['blocks'].append({'type':'table','rows':rows})
    return {'filename':path.name,'sections':sections,'paragraph_count':len(doc.paragraphs),'table_count':len(doc.tables),'status':'Archived source report; not a re-audited or current investment recommendation.'}

def xlsx_xml(path):
    # Direct OOXML extraction preserves original cached values and formulas exactly.
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    z=zipfile.ZipFile(path); strings=[]
    if 'xl/sharedStrings.xml' in z.namelist():
        ss=ET.fromstring(z.read('xl/sharedStrings.xml'))
        strings=[''.join(s.itertext()) for s in ss]
    rels={r.attrib['Id']:r.attrib['Target'] for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    book=ET.fromstring(z.read('xl/workbook.xml')); sheets=[]
    for sn in book.find('m:sheets',ns):
        name=sn.attrib['name']; target=rels[sn.attrib['{'+ns['r']+'}id']]
        target=target.lstrip('/') if target.startswith('/') else 'xl/'+target
        sheet=ET.fromstring(z.read(target)); cells=[]; maxr=maxc=0
        for c in sheet.findall('.//m:sheetData/m:row/m:c',ns):
            a=c.attrib['r']; row=int(re.sub('[A-Z]','',a)); col=0
            for ch in re.sub('[0-9]','',a): col=col*26+ord(ch)-64
            val=c.find('m:v',ns); form=c.find('m:f',ns); typ=c.attrib.get('t')
            if typ=='inlineStr':v=''.join(c.find('m:is',ns).itertext())
            elif val is None:v=None
            elif typ=='s':v=strings[int(val.text)]
            else:
                v=val.text
                if typ not in ('str','e') and v is not None:
                    try:v=float(v);v=int(v) if v.is_integer() else v
                    except ValueError:pass
            if v is None and form is None:continue
            cells.append({'address':a,'row':row,'col':col,'value':v,'formula':'='+form.text if form is not None else None})
            maxr=max(maxr,row);maxc=max(maxc,col)
        sheets.append({'name':name,'cells':cells,'max_row':maxr,'max_col':maxc})
    return {'filename':path.name,'sheets':sheets,'status':'Archived attachment; formulas and original values preserved, not silently recalculated.'}

reports=[docx_blocks(P/'compute_infrastructure_report_per_gw_costs.docx'),docx_blocks(P/'compute_infrastructure_report.docx')]
workbooks=[xlsx_xml(P/'compute_infrastructure_report_per_gw_costs.xlsx'),xlsx_xml(P/'compute_infrastructure_report.xlsx')]
legacy=[json.load(open(P/'compute_model_data_per_gw_costs.json')),json.load(open(P/'compute_model_data.json'))]
archive={'reports':reports,'workbooks':workbooks,'legacy':legacy}
(OUT/'archive.json').write_text(json.dumps(archive,ensure_ascii=False,separators=(',',':')))
manifest=[]
for name in ['compute_infrastructure_report_per_gw_costs.pdf','compute_infrastructure_report.pdf','compute_infrastructure_report_per_gw_costs.xlsx','compute_infrastructure_report.xlsx','compute_infrastructure_report_per_gw_costs.docx','compute_infrastructure_report.docx','compute_model_data_per_gw_costs.json','compute_model_data.json']:
 p=P/name; manifest.append({'filename':name,'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest(),'status':'Archived original'})
(OUT/'manifest.json').write_text(json.dumps(manifest,indent=2))
print('Sections:',[len(r['sections']) for r in reports],'Tables:',[r['table_count'] for r in reports], 'Workbook cells:',[sum(len(s['cells']) for s in w['sheets']) for w in workbooks])

# The distributed world geometry is already available. Rebuild only if absent.
if (OUT/'world.json').exists():
 print('Keeping the bundled Natural Earth geometry.')
 raise SystemExit(0)
# Natural Earth low-resolution country outlines and equal-area-ish land points.
import pyogrio, shapely, numpy as np
f=Path('/opt/pyvenv/lib/python3.13/site-packages/pyogrio/tests/fixtures/naturalearth_lowres/naturalearth_lowres.shp')
gdf=pyogrio.read_dataframe(f); lines=[]
for geom in gdf.geometry:
 geom=geom.simplify(.16)
 for pol in (geom.geoms if geom.geom_type=='MultiPolygon' else [geom]):
  points=list(pol.exterior.coords); line=[]
  for a,b in zip(points,points[1:]):
   dist=max(abs(a[0]-b[0]),abs(a[1]-b[1])); n=max(1,int(dist/1.8))
   for k in range(n): line.append([round(a[0]+(b[0]-a[0])*k/n,3),round(a[1]+(b[1]-a[1])*k/n,3)])
  line.append([round(x,3) for x in points[-1]]);lines.append(line)
land=shapely.union_all(gdf.geometry.values);dots=[]
for lat in np.arange(-82,83,1.65):
 step=1.65/max(.15,np.cos(np.radians(lat)))
 for lon in np.arange(-180,180,step):
  if land.contains(shapely.Point(lon,lat)):dots.append([round(lon,3),round(lat,3)])
world={'lines':lines,'dots':dots,'attribution':'Natural Earth (public domain), low-resolution geographic context; boundaries are illustrative and not an endorsement.'}
(OUT/'world.json').write_text(json.dumps(world,separators=(',',':')))
print('Land dots',len(dots),'lines',len(lines),'archive MB',(OUT/'archive.json').stat().st_size/1e6)
